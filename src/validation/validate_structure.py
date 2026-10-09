"""
Structural validation.

Checks that each ingested file can be read and has the shape its
source contract (config/source_registry.yaml) promises:

    xlsx  workbook readable, required sheet, required columns,
          duplicate / unexpected columns, minimum rows
    csv   file readable, required columns, duplicate / unexpected
          columns, minimum rows
    pdf   document readable, file-naming convention, required labels

Usage:
    python src/validation/validate_structure.py
    python src/validation/validate_structure.py --run-id ING-20261004_154838
    python src/validation/validate_structure.py --source supplier_master
    python src/validation/validate_structure.py --source supplier_master \
        --file tests/test_data/corrupted_supplier_master.xlsx
"""

import argparse
import csv
import re
import sys

from collections import Counter
from itertools import islice
from pathlib import Path

import openpyxl

from pypdf import PdfReader

from common import (
    PROJECT_ROOT,
    is_blank,
    load_source_registry,
    make_result,
    quarantine_file,
    select_targets,
    write_file_status,
    write_validation_results,
)


STAGE = "STRUCTURE_VALIDATION"


# ---------------------------------------------------------
# FILE INSPECTION
# ---------------------------------------------------------

def count_data_rows(
    rows,
    row_limit: int | None,
) -> int:
    """
    Count non-blank rows. With a row_limit, stop as soon as that many
    have been seen - enough to prove a minimum without reading a
    large file to the end.
    """

    non_blank = (
        row for row in rows
        if not all(is_blank(value) for value in row)
    )

    return sum(1 for _ in islice(non_blank, row_limit))


def inspect_workbook(
    file: Path,
    sheet_name: str | None,
    row_limit: int | None = None,
) -> dict:
    """
    Open a workbook once and collect what the structural checks
    need. Raises if the file is not a readable workbook.
    """

    workbook = openpyxl.load_workbook(
        file,
        read_only=True,
        data_only=True,
    )

    try:

        sheet_names = workbook.sheetnames

        if sheet_name not in sheet_names:
            return {
                "sheets": sheet_names,
                "headers": None,
                "row_count": None,
            }

        rows = workbook[sheet_name].iter_rows(values_only=True)

        headers = [
            header.strip() if isinstance(header, str) else header
            for header in next(rows, ())
        ]

        row_count = count_data_rows(rows, row_limit)

    finally:
        workbook.close()

    return {
        "sheets": sheet_names,
        "headers": headers,
        "row_count": row_count,
    }


def inspect_csv(
    file: Path,
    row_limit: int | None = None,
) -> dict:

    with open(
        file,
        "r",
        newline="",
        encoding="utf-8-sig",
    ) as csv_file:

        reader = csv.reader(csv_file)

        headers = [header.strip() for header in next(reader, [])]

        row_count = count_data_rows(reader, row_limit)

    return {
        "headers": headers,
        "row_count": row_count,
    }


# ---------------------------------------------------------
# TABULAR CHECKS
# ---------------------------------------------------------

def validate_readable(
    file: Path,
    run_id: str,
    source: str,
    structure: dict,
    row_limit: int | None = None,
) -> tuple[dict, dict | None]:
    """
    Check whether the file can be opened. Returns the validation
    result and, if readable, the inspection details.
    """

    is_excel = file.suffix.lower() == ".xlsx"

    rule = "WORKBOOK_READABLE" if is_excel else "FILE_READABLE"

    try:

        inspection = (
            inspect_workbook(file, structure.get("sheet"), row_limit)
            if is_excel
            else inspect_csv(file, row_limit)
        )

    except Exception as error:

        kind = "an Excel workbook" if is_excel else "a UTF-8 CSV file"

        return make_result(
            run_id=run_id,
            source=source,
            file=file,
            stage=STAGE,
            rule=rule,
            status="FAIL",
            severity="ERROR",
            error_type=type(error).__name__,
            technical_error=str(error),
            expected=f"A valid readable {kind}.",
            message=f"The file could not be opened as {kind}.",
            recommended_action=(
                "Verify that the file is genuine, is not corrupted, "
                "and was not renamed from another file type. "
                "Re-export it from the source system and resubmit."
            ),
        ), None

    return make_result(
        run_id=run_id,
        source=source,
        file=file,
        stage=STAGE,
        rule=rule,
        status="PASS",
        message="File opened successfully.",
        actual=(
            f"Sheets found: {inspection['sheets']}"
            if is_excel else None
        ),
    ), inspection


def validate_required_sheet(
    file: Path,
    run_id: str,
    source: str,
    structure: dict,
    inspection: dict,
) -> dict:

    expected_sheet = structure["sheet"]
    actual_sheets = inspection["sheets"]

    if expected_sheet in actual_sheets:

        return make_result(
            run_id=run_id,
            source=source,
            file=file,
            sheet=expected_sheet,
            stage=STAGE,
            rule="REQUIRED_SHEET",
            status="PASS",
            actual=actual_sheets,
            expected=expected_sheet,
            message=f"Required worksheet '{expected_sheet}' exists.",
        )

    return make_result(
        run_id=run_id,
        source=source,
        file=file,
        stage=STAGE,
        rule="REQUIRED_SHEET",
        status="FAIL",
        severity="ERROR",
        actual=actual_sheets,
        expected=expected_sheet,
        message=f"Required worksheet '{expected_sheet}' is missing.",
        recommended_action=(
            f"Add or restore worksheet '{expected_sheet}' "
            f"and resubmit the file."
        ),
    )


def validate_required_columns(
    file: Path,
    run_id: str,
    source: str,
    structure: dict,
    inspection: dict,
) -> list[dict]:

    results = []

    sheet = structure.get("sheet")
    required_columns = structure.get("required_columns") or []
    actual_columns = inspection["headers"]

    for column in required_columns:

        if column in actual_columns:

            results.append(make_result(
                run_id=run_id,
                source=source,
                file=file,
                sheet=sheet,
                row_number=1,
                stage=STAGE,
                rule="REQUIRED_COLUMN",
                field=column,
                status="PASS",
                actual=column,
                expected=column,
                message=f"Required column '{column}' exists.",
            ))

        else:

            results.append(make_result(
                run_id=run_id,
                source=source,
                file=file,
                sheet=sheet,
                row_number=1,
                stage=STAGE,
                rule="REQUIRED_COLUMN",
                field=column,
                status="FAIL",
                severity="ERROR",
                actual="COLUMN_NOT_FOUND",
                expected=column,
                message=f"Required column '{column}' is missing.",
                recommended_action=(
                    f"Add or restore the required column "
                    f"'{column}' in the header row and resubmit "
                    f"the file."
                ),
            ))

    return results


def validate_column_layout(
    file: Path,
    run_id: str,
    source: str,
    structure: dict,
    inspection: dict,
) -> list[dict]:
    """
    Duplicate column names make field lookups ambiguous (ERROR).
    Columns the contract does not know about indicate schema drift
    (WARNING) - the data is still usable.
    """

    results = []

    sheet = structure.get("sheet")
    required_columns = structure.get("required_columns") or []
    headers = inspection["headers"]

    named_headers = [header for header in headers if not is_blank(header)]

    for column, count in Counter(named_headers).items():

        if count > 1:

            results.append(make_result(
                run_id=run_id,
                source=source,
                file=file,
                sheet=sheet,
                row_number=1,
                stage=STAGE,
                rule="DUPLICATE_COLUMN",
                field=column,
                status="FAIL",
                severity="ERROR",
                actual=f"{count} columns named '{column}'",
                expected="Each column name appears once.",
                message=f"Column '{column}' appears {count} times.",
                recommended_action=(
                    f"Remove or rename the duplicate '{column}' "
                    f"columns and resubmit the file."
                ),
            ))

    if required_columns:

        for column in dict.fromkeys(named_headers):

            if column not in required_columns:

                results.append(make_result(
                    run_id=run_id,
                    source=source,
                    file=file,
                    sheet=sheet,
                    row_number=1,
                    stage=STAGE,
                    rule="UNEXPECTED_COLUMN",
                    field=column,
                    status="FAIL",
                    severity="WARNING",
                    actual=column,
                    expected=f"One of {required_columns}",
                    message=(
                        f"Column '{column}' is not defined in the "
                        f"source contract."
                    ),
                    recommended_action=(
                        "Confirm whether the source layout changed. "
                        "If the column is legitimate, add it to the "
                        "source registry."
                    ),
                ))

    blank_positions = [
        position
        for position, header in enumerate(headers, start=1)
        if is_blank(header)
    ]

    if blank_positions:

        results.append(make_result(
            run_id=run_id,
            source=source,
            file=file,
            sheet=sheet,
            row_number=1,
            stage=STAGE,
            rule="BLANK_COLUMN_HEADER",
            status="FAIL",
            severity="WARNING",
            actual=f"Blank header in column position(s) {blank_positions}",
            expected="Every column has a header.",
            message="One or more columns have no header.",
            recommended_action=(
                "Name or remove the unnamed columns. Values in them "
                "are ignored by validation."
            ),
        ))

    return results


def validate_minimum_rows(
    file: Path,
    run_id: str,
    source: str,
    structure: dict,
    inspection: dict,
) -> dict:

    minimum_rows = structure.get("minimum_rows", 1)
    row_count = inspection["row_count"]

    if row_count >= minimum_rows:

        return make_result(
            run_id=run_id,
            source=source,
            file=file,
            sheet=structure.get("sheet"),
            stage=STAGE,
            rule="MINIMUM_ROWS",
            status="PASS",
            actual=row_count,
            expected=f">= {minimum_rows}",
            message=f"File contains {row_count} data rows.",
        )

    return make_result(
        run_id=run_id,
        source=source,
        file=file,
        sheet=structure.get("sheet"),
        stage=STAGE,
        rule="MINIMUM_ROWS",
        status="FAIL",
        severity="ERROR",
        actual=row_count,
        expected=f">= {minimum_rows}",
        message=(
            f"File contains {row_count} data rows; "
            f"at least {minimum_rows} expected."
        ),
        recommended_action=(
            "Confirm the extract was not truncated or exported "
            "empty, then resubmit the complete file."
        ),
    )


# ---------------------------------------------------------
# PDF CHECKS
# ---------------------------------------------------------

def validate_pdf_document(
    file: Path,
    run_id: str,
    source: str,
    structure: dict,
) -> list[dict]:

    results = []

    try:

        reader = PdfReader(file)

        text = "\n".join(
            page.extract_text() or ""
            for page in reader.pages
        )

    except Exception as error:

        return [make_result(
            run_id=run_id,
            source=source,
            file=file,
            stage=STAGE,
            rule="PDF_READABLE",
            status="FAIL",
            severity="ERROR",
            error_type=type(error).__name__,
            technical_error=str(error),
            expected="A valid readable PDF document.",
            message="The file could not be opened as a PDF document.",
            recommended_action=(
                "Verify that the PDF is not corrupted or password "
                "protected, then resubmit it."
            ),
        )]

    if not text.strip():

        return [make_result(
            run_id=run_id,
            source=source,
            file=file,
            stage=STAGE,
            rule="PDF_READABLE",
            status="FAIL",
            severity="ERROR",
            actual=f"{len(reader.pages)} page(s), no extractable text",
            expected="A PDF with a text layer.",
            message="The PDF contains no extractable text (possibly a scan).",
            recommended_action=(
                "Submit the system-generated PDF rather than a "
                "scanned copy."
            ),
        )]

    results.append(make_result(
        run_id=run_id,
        source=source,
        file=file,
        stage=STAGE,
        rule="PDF_READABLE",
        status="PASS",
        actual=f"{len(reader.pages)} page(s)",
        message="PDF opened and text extracted successfully.",
    ))

    filename_pattern = structure.get("filename_pattern")

    if filename_pattern:

        matches = re.fullmatch(filename_pattern, file.name) is not None

        results.append(make_result(
            run_id=run_id,
            source=source,
            file=file,
            stage=STAGE,
            rule="FILENAME_CONVENTION",
            status="PASS" if matches else "FAIL",
            severity=None if matches else "ERROR",
            actual=file.name,
            expected=filename_pattern,
            message=(
                "File name follows the naming convention."
                if matches else
                "File name does not follow the naming convention."
            ),
            recommended_action=(
                None if matches else
                "Rename the file to <SiteCode>_<ShiftID>_<YYYYMMDD>"
                "_Cashup.pdf and resubmit."
            ),
        ))

    for label in structure.get("required_labels") or []:

        found = f"{label}:" in text

        results.append(make_result(
            run_id=run_id,
            source=source,
            file=file,
            stage=STAGE,
            rule="REQUIRED_LABEL",
            field=label,
            status="PASS" if found else "FAIL",
            severity=None if found else "ERROR",
            actual=None if found else "LABEL_NOT_FOUND",
            expected=f"{label}:",
            message=(
                f"Label '{label}' found."
                if found else
                f"Required label '{label}' is missing from the document."
            ),
            recommended_action=(
                None if found else
                "Regenerate the document from the standard template "
                "and resubmit."
            ),
        ))

    return results


# ---------------------------------------------------------
# SOURCE VALIDATION
# ---------------------------------------------------------

def first_failed_rule(results: list[dict]) -> str | None:

    for result in results:

        if (
            result["status"] == "FAIL"
            and result.get("severity") == "ERROR"
        ):
            return result["rule"]

    return None


def summarise(
    file: Path,
    run_id: str,
    source: str,
    results: list[dict],
    row_count: int | None = None,
) -> dict:

    failed_rule = first_failed_rule(results)

    return {
        "run_id": run_id,
        "source": source,
        "file": file.name,
        "path": file,
        "status": "STRUCTURE_INVALID" if failed_rule else "STRUCTURE_VALID",
        "failed_stage": failed_rule,
        "row_count": row_count,
        "results": results,
    }


def validate_source(
    file: Path,
    run_id: str,
    source: str,
    source_registry: dict,
    full_row_count: bool = True,
) -> dict:
    """
    Run all structural checks for one file. With full_row_count=False
    the row count stops at the contract's minimum_rows, which is all
    a pass/fail gate needs.
    """

    source_config = source_registry[source]
    structure = source_config.get("structure") or {}
    file_format = source_config["format"].lower()

    if file_format == "pdf":

        results = validate_pdf_document(
            file=file,
            run_id=run_id,
            source=source,
            structure=structure,
        )

        return summarise(file, run_id, source, results)

    validation_results = []

    # -----------------------------------------------------
    # 1. FILE READABILITY
    # -----------------------------------------------------

    readable_result, inspection = validate_readable(
        file=file,
        run_id=run_id,
        source=source,
        structure=structure,
        row_limit=(
            None if full_row_count
            else structure.get("minimum_rows", 1)
        ),
    )

    validation_results.append(readable_result)

    if inspection is None:
        return summarise(file, run_id, source, validation_results)

    # -----------------------------------------------------
    # 2. REQUIRED SHEET (Excel only)
    # -----------------------------------------------------

    if file_format == "xlsx":

        sheet_result = validate_required_sheet(
            file=file,
            run_id=run_id,
            source=source,
            structure=structure,
            inspection=inspection,
        )

        validation_results.append(sheet_result)

        if sheet_result["status"] == "FAIL":
            return summarise(file, run_id, source, validation_results)

    # -----------------------------------------------------
    # 3. COLUMNS
    # -----------------------------------------------------

    validation_results.extend(validate_required_columns(
        file=file,
        run_id=run_id,
        source=source,
        structure=structure,
        inspection=inspection,
    ))

    validation_results.extend(validate_column_layout(
        file=file,
        run_id=run_id,
        source=source,
        structure=structure,
        inspection=inspection,
    ))

    # -----------------------------------------------------
    # 4. MINIMUM ROWS
    # -----------------------------------------------------

    validation_results.append(validate_minimum_rows(
        file=file,
        run_id=run_id,
        source=source,
        structure=structure,
        inspection=inspection,
    ))

    return summarise(
        file,
        run_id,
        source,
        validation_results,
        row_count=inspection["row_count"],
    )


# ---------------------------------------------------------
# COMMAND LINE
# ---------------------------------------------------------

def parse_arguments(description: str) -> argparse.Namespace:

    parser = argparse.ArgumentParser(description=description)

    parser.add_argument(
        "--run-id",
        help="Ingestion run to validate (default: latest run that ingested files).",
    )
    parser.add_argument(
        "--source",
        help="Only validate this registered source.",
    )
    parser.add_argument(
        "--file",
        help="Validate a file outside the pipeline (requires --source).",
    )
    parser.add_argument(
        "--no-quarantine",
        action="store_true",
        help="Do not copy failed files to data/quarantine.",
    )

    return parser.parse_args()


def main() -> int:

    args = parse_arguments(
        "Validate the structure of ingested source files."
    )

    source_registry = load_source_registry()

    run_id, targets, is_adhoc = select_targets(
        args.run_id,
        args.source,
        args.file,
        source_registry,
    )

    quarantine_enabled = not (is_adhoc or args.no_quarantine)

    print(f"\nSTRUCTURE VALIDATION: {run_id}")
    print(f"Files: {len(targets)}\n")

    status_counts = Counter()

    for source, file in targets:

        validation = validate_source(
            file=file,
            run_id=run_id,
            source=source,
            source_registry=source_registry,
        )

        results = validation["results"]

        write_validation_results(results)

        quarantine_path = None

        if (
            validation["status"] == "STRUCTURE_INVALID"
            and quarantine_enabled
        ):
            quarantine_path = quarantine_file(file, run_id, results)

        write_file_status(
            run_id=run_id,
            source=source,
            file=file,
            stage=STAGE,
            status=validation["status"],
            results=results,
            rows_checked=validation["row_count"],
            quarantine_path=quarantine_path,
        )

        status_counts[validation["status"]] += 1

        warnings = sum(
            1 for result in results
            if result.get("severity") == "WARNING"
        )

        line = f"{file.name} -> {source} -> {validation['status']}"

        if validation["failed_stage"]:
            line += f" (failed: {validation['failed_stage']})"

        if warnings:
            line += f" [{warnings} warning(s)]"

        print(line)

        for result in results:

            if result["status"] == "FAIL":
                print(
                    f"   {result['severity']}: {result['rule']}"
                    f" -> {result['message']}"
                )

        if quarantine_path:
            print(f"   Quarantined: {quarantine_path.relative_to(PROJECT_ROOT)}")

    print("\nVALIDATION SUMMARY")
    print("--------------------------------------------")
    print(f"Run ID:            {run_id}")
    print(f"Files validated:   {len(targets)}")
    print(f"Structure valid:   {status_counts['STRUCTURE_VALID']}")
    print(f"Structure invalid: {status_counts['STRUCTURE_INVALID']}")

    return 1 if status_counts["STRUCTURE_INVALID"] else 0


if __name__ == "__main__":

    sys.exit(main())
