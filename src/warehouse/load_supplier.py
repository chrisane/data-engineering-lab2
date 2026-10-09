"""Controlled Supplier Excel -> SQL Server staging loader.

Run from the project root:
    python src/warehouse/load_supplier.py

Requires existing SQL objects: stg.Supplier and audit.LoadExecution.
Uses the project's authoritative eligibility decision logic; never loads incoming/.
"""

import hashlib
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "src" / "validation"))

from common import (  # noqa: E402
    FILE_STATUS_LOG_FILE,
    VALIDATION_LOG_FILE,
    load_registry_config,
    read_log,
    read_manifest,
)
from load_eligibility import decide  # noqa: E402
from db_connection import get_connection  # noqa: E402

SOURCE = "supplier_master"
COLUMNS = (
    "supplier_id", "supplier_name", "location",
    "minimum_order_value_nad", "rebate_rate", "payment_terms_days",
)


def source_rows(raw_path: Path, sheet_name: str):
    """Read a validated workbook and retain original Excel row numbers."""
    workbook = openpyxl.load_workbook(raw_path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Required sheet missing: {sheet_name}")
        sheet = workbook[sheet_name]
        iterator = sheet.iter_rows(values_only=True)
        header_row = next(iterator, None)
        if header_row is None:
            raise ValueError("Supplier workbook is empty")
        headers = [str(v).strip() if v is not None else "" for v in header_row]
        if len(headers) != len(set(headers)):
            raise ValueError("Duplicate column headings in supplier workbook")
        missing = set(COLUMNS) - set(headers)
        if missing:
            raise ValueError(f"Missing supplier columns: {sorted(missing)}")
        positions = [headers.index(name) for name in COLUMNS]
        result = []
        for row_number, cells in enumerate(iterator, start=2):
            if not any(v is not None and str(v).strip() for v in cells):
                continue
            values = [cells[i] if i < len(cells) else None for i in positions]
            result.append((row_number, *values))
        return result
    finally:
        workbook.close()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def audit_attempt(connection, decision, status, rows_read=0, rows_inserted=0,
                  rows_rejected=0, error=None):
    """Persist a terminal audit record in a separate committed transaction."""
    cursor = connection.cursor()
    cursor.execute(
        """
        INSERT INTO audit.LoadExecution
            (IngestionRunID, SourceName, SourceFile, SourceFileHash,
             LoadStatus, RowsRead, RowsInserted, RowsRejected,
             CompletedAt, ErrorMessage)
        OUTPUT INSERTED.LoadID
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, SYSUTCDATETIME(), ?)
        """,
        decision.get("run_id") or "UNKNOWN",
        SOURCE,
        decision.get("file") or "UNKNOWN",
        (decision.get("file_hash") or "0" * 64).lower(),
        status,
        rows_read,
        rows_inserted,
        rows_rejected,
        str(error)[:4000] if error else None,
    )
    load_id = cursor.fetchone()[0]
    connection.commit()
    return load_id


def get_supplier_decision():
    config = load_registry_config()
    decisions = decide(
        config,
        read_manifest(),
        read_log(FILE_STATUS_LOG_FILE),
        read_log(VALIDATION_LOG_FILE),
    )
    matches = [row for row in decisions if row["source"] == SOURCE]
    if len(matches) != 1:
        raise RuntimeError(
            f"Expected exactly one {SOURCE} decision; found {len(matches)}"
        )
    return config, matches[0]


def main() -> int:
    config, decision = get_supplier_decision()
    print(f"Source: {SOURCE}")
    print(f"Decision: {decision['decision']}")
    print(f"Ingestion run: {decision.get('run_id')}")

    # The gate is evaluated fresh, not taken from a potentially stale CSV decision.
    if decision["decision"] != "ELIGIBLE":
        reason = decision.get("reason") or "Load not authorized"
        print(f"SQL LOAD REFUSED: {reason}")
        if decision.get("run_id") and decision.get("file_hash"):
            connection = get_connection()
            try:
                audit_attempt(connection, decision, "SKIPPED", error=reason)
            finally:
                connection.close()
        return 2

    raw = Path(decision["raw_path"])
    if not raw.is_absolute():
        raw = ROOT / raw
    raw = raw.resolve()
    raw_root = (ROOT / "data" / "raw").resolve()

    # Reject a source path outside immutable raw storage.
    if not raw.is_relative_to(raw_root) or not raw.is_file():
        raise RuntimeError(f"Raw file missing or outside data/raw: {raw}")

    digest = sha256_file(raw)
    if digest.lower() != decision["file_hash"].lower():
        raise RuntimeError("Raw SHA-256 mismatch: file has changed since ingestion")

    sheet_name = config["sources"][SOURCE]["structure"]["sheet"]
    rows = source_rows(raw, sheet_name)
    if not rows:
        raise RuntimeError("Eligible Supplier workbook contains no data rows")

    connection = get_connection()
    connection.autocommit = False
    lock_acquired = False
    try:
        cursor = connection.cursor()
        # Serialize Supplier loads, including concurrent executions of this script.
        cursor.execute(
            """
            DECLARE @result int;
            EXEC @result = sys.sp_getapplock
                @Resource = 'RetailDataPlatform:stg.Supplier:load',
                @LockMode = 'Exclusive',
                @LockOwner = 'Session',
                @LockTimeout = 10000;
            SELECT @result;
            """
        )
        lock_result = cursor.fetchone()[0]
        if lock_result < 0:
            raise RuntimeError(f"Could not acquire Supplier load lock ({lock_result})")
        lock_acquired = True

        # Existing staging data must be complete; a partial prior load is an error.
        cursor.execute(
            """
            SELECT COUNT(*) FROM stg.Supplier
            WHERE IngestionRunID = ? AND SourceFileHash = ?
            """,
            decision["run_id"], digest,
        )
        existing = cursor.fetchone()[0]
        if existing:
            if existing != len(rows):
                raise RuntimeError(
                    f"Existing staging count {existing} differs from source {len(rows)}; "
                    "manual reconciliation required"
                )
            # Check row-number completeness, not just count.
            cursor.execute(
                """
                SELECT SourceRowNumber FROM stg.Supplier
                WHERE IngestionRunID = ? AND SourceFileHash = ?
                """,
                decision["run_id"], digest,
            )
            staged_numbers = {record[0] for record in cursor.fetchall()}
            if staged_numbers != {record[0] for record in rows}:
                raise RuntimeError("Existing staged row numbers differ from source")
            connection.rollback()
            load_id = audit_attempt(
                connection, decision, "SKIPPED", rows_read=len(rows),
                error="Already loaded; row counts and source row numbers match",
            )
            print(f"SKIPPED_DUPLICATE: {len(rows)} rows already staged; audit LoadID={load_id}")
            return 0

        # A single transaction protects all staging inserts and the COMPLETED audit.
        cursor.execute(
            """
            INSERT INTO audit.LoadExecution
                (IngestionRunID, SourceName, SourceFile, SourceFileHash,
                 LoadStatus, RowsRead)
            OUTPUT INSERTED.LoadID
            VALUES (?, ?, ?, ?, 'STARTED', ?)
            """,
            decision["run_id"], SOURCE, decision["file"], digest, len(rows),
        )
        load_id = cursor.fetchone()[0]

        insert_sql = """
            INSERT INTO stg.Supplier
                (supplier_id, supplier_name, location, minimum_order_value_nad,
                 rebate_rate, payment_terms_days, IngestionRunID, SourceFile,
                 SourceRowNumber, SourceFileHash)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """
        for row_number, supplier_id, supplier_name, location, minimum, rebate, terms in rows:
            cursor.execute(
                insert_sql,
                supplier_id, supplier_name, location, minimum, rebate, terms,
                decision["run_id"], decision["file"], row_number, digest,
            )

        cursor.execute(
            """
            SELECT COUNT(*) FROM stg.Supplier
            WHERE IngestionRunID = ? AND SourceFileHash = ?
            """,
            decision["run_id"], digest,
        )
        staged_count = cursor.fetchone()[0]
        if staged_count != len(rows):
            raise RuntimeError(
                f"Row reconciliation failed: source={len(rows)}, SQL={staged_count}"
            )

        cursor.execute(
            """
            UPDATE audit.LoadExecution
            SET LoadStatus='COMPLETED', RowsInserted=?,
                CompletedAt=SYSUTCDATETIME()
            WHERE LoadID=?
            """,
            staged_count, load_id,
        )
        connection.commit()
        print(f"SQL LOAD COMPLETED: {staged_count} rows; audit LoadID={load_id}")
        return 0

    except Exception as exc:
        connection.rollback()
        try:
            load_id = audit_attempt(
                connection, decision, "FAILED", rows_read=len(rows), error=exc,
            )
            print(f"SQL LOAD FAILED: {exc}; audit LoadID={load_id}")
        except Exception as audit_exc:
            print(f"SQL LOAD FAILED: {exc}")
            print(f"WARNING: Could not record failure audit: {audit_exc}")
        return 1
    finally:
        if lock_acquired:
            try:
                connection.cursor().execute(
                    "EXEC sys.sp_releaseapplock "
                    "@Resource='RetailDataPlatform:stg.Supplier:load', "
                    "@LockOwner='Session'"
                )
            except Exception:
                pass
        connection.close()


if __name__ == "__main__":
    sys.exit(main())
